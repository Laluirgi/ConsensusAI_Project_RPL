import sys

path = r'd:\AIKB\backend\main.py'
with open(path, 'r', encoding='utf-8') as f:
    code = f.read()

old_plan = '''@app.post("/plan")
async def plan_endpoint(request: Request):
    data = await request.json()
    question = data.get("question", "")
    system_prompt = "Kamu adalah Planner AI. Jika input user adalah sapaan ringan, percakapan santai (seperti 'halo', 'apa kabar'), atau sekadar basa-basi yang tidak butuh rencana/debat rumit, maka jawablah langsung layaknya teman mengobrol yang asik, dan WAJIB tambahkan string '[SKIP_DEBATE]' persis di akhir jawabanmu. JIKA input user adalah permintaan tugas (coding, analisis, buatan soal, dll), Buatkan rencana strategis, komprehensif, dan terstruktur langkah demi langkah untuk menjawabnya tanpa string tersebut."
    
    return StreamingResponse(
        generate_stream("Gemini Flash Core", system_prompt, f"Permintaan user: {question}"),
        media_type="text/event-stream"
    )'''

new_plan = '''@app.post("/plan")
async def plan_endpoint(request: Request):
    data = await request.json()
    question = data.get("question", "")
    model_name = data.get("model", "Gemini Flash Core")
    system_prompt = "Kamu adalah Planner AI. Jika input user adalah sapaan ringan, percakapan santai (seperti 'halo', 'apa kabar'), atau sekadar basa-basi yang tidak butuh rencana/debat rumit, maka jawablah langsung layaknya teman mengobrol yang asik, dan WAJIB tambahkan string '[SKIP_DEBATE]' persis di akhir jawabanmu. JIKA input user adalah permintaan tugas (coding, analisis, buatan soal, dll), Buatkan rencana strategis, komprehensif, dan terstruktur langkah demi langkah untuk menjawabnya tanpa string tersebut."
    
    return StreamingResponse(
        generate_stream(model_name, system_prompt, f"Permintaan user: {question}"),
        media_type="text/event-stream"
    )'''

code = code.replace(old_plan, new_plan)

export_code = '''
import io
try:
    from docx import Document
except ImportError:
    pass
try:
    from openpyxl import Workbook
except ImportError:
    pass
from fastapi.responses import Response

@app.post("/export")
async def export_endpoint(request: Request):
    data = await request.json()
    content = data.get("content", "")
    format_type = data.get("format", "docx")
    
    if format_type == "docx":
        doc = Document()
        doc.add_heading("Consensus Result", 0)
        for line in content.split("\\n"):
            doc.add_paragraph(line)
        
        f = io.BytesIO()
        doc.save(f)
        f.seek(0)
        return Response(content=f.read(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=result.docx"})
        
    elif format_type == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Result"
        
        for i, line in enumerate(content.split("\\n"), 1):
            if "|" in line:
                row = [x.strip() for x in line.split("|") if x.strip()]
                if row:
                    ws.append(row)
            else:
                ws.cell(row=i, column=1, value=line)
                
        f = io.BytesIO()
        wb.save(f)
        f.seek(0)
        return Response(content=f.read(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=result.xlsx"})
        
    return {"error": "Invalid format"}

'''

code = code.replace('if __name__ == "__main__":', export_code + 'if __name__ == "__main__":')

with open(path, 'w', encoding='utf-8') as f:
    f.write(code)
print('Patched main.py!')
