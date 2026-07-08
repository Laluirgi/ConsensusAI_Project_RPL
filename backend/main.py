from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from contextlib import asynccontextmanager
import uvicorn
import io
import json

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    pass
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
except ImportError:
    pass

from agents import generate_stream
from database import (
    init_db, add_feedback, get_all_feedbacks, update_feedback_status,
    delete_feedback, add_rating, get_all_ratings, get_rating_stats,
    get_feedback_stats
)


# ── App Lifecycle ──
@asynccontextmanager
async def lifespan(app: FastAPI):
    """Initialize database on startup."""
    await init_db()
    yield


app = FastAPI(title="Consensus AI Backend", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ══════════════════════════════════════════
# AI Endpoints
# ══════════════════════════════════════════

@app.post("/plan")
async def plan_endpoint(request: Request):
    """Fase 1: Planner membuat rencana jawaban."""
    data = await request.json()
    question = data.get("question", "")
    model_name = data.get("model", "Gemini Flash")
    files = data.get("files", [])
    system_prompt = (
        "Kamu adalah Planner AI. "
        "Jika input user adalah sapaan ringan, percakapan santai (seperti 'halo', 'apa kabar'), "
        "atau sekadar basa-basi yang tidak butuh rencana/debat rumit, maka jawablah langsung "
        "layaknya teman mengobrol yang asik, dan WAJIB tambahkan string '[SKIP_DEBATE]' persis "
        "di akhir jawabanmu. JIKA input user adalah permintaan tugas (coding, analisis, buatan soal, dll), "
        "Buatkan rencana strategis, komprehensif, dan terstruktur langkah demi langkah untuk "
        "menjawabnya tanpa string tersebut."
    )
    
    return StreamingResponse(
        generate_stream(model_name, system_prompt, f"Permintaan user: {question}", files),
        media_type="text/event-stream"
    )


@app.post("/solo")
async def solo_endpoint(request: Request):
    """Mode Solo: satu AI menjawab langsung tanpa debat."""
    data = await request.json()
    question = data.get("question", "")
    model_name = data.get("model", "Gemini Flash")
    files = data.get("files", [])
    history = data.get("history", [])
    
    # Build system prompt with conversation history for follow-up
    history_context = ""
    if history:
        history_context = "\n\nRiwayat percakapan sebelumnya:\n"
        for h in history:
            role = "User" if h.get("role") == "user" else "AI"
            history_context += f"{role}: {h.get('content', '')}\n\n"
    
    system_prompt = (
        f"Kamu adalah {model_name}, asisten AI yang cerdas dan membantu. "
        "Jawab pertanyaan user secara komprehensif, detail, dan terstruktur. "
        "Gunakan heading, bullet points, dan formatting markdown yang rapi. "
        "Jika ada kode, berikan kode lengkap yang bisa langsung dijalankan. "
        "Sertakan alasan di setiap poin penting."
        f"{history_context}"
    )
    
    return StreamingResponse(
        generate_stream(model_name, system_prompt, question, files),
        media_type="text/event-stream"
    )


@app.post("/debate")
async def debate_endpoint(request: Request):
    """Fase 2: Debat/diskusi antar AI."""
    data = await request.json()
    system_prompt = data.get("system_prompt", "")
    message = data.get("message", "")
    model_name = data.get("model", "Gemini Flash")
    files = data.get("files", [])
    
    return StreamingResponse(
        generate_stream(model_name, system_prompt, message, files),
        media_type="text/event-stream"
    )


@app.post("/revise")
async def revise_endpoint(request: Request):
    """Revisi jawaban tanpa debat ulang."""
    data = await request.json()
    revision_request = data.get("revision_request", "")
    original_answer = data.get("original_answer", "")
    model_name = data.get("model", "Gemini Flash")
    files = data.get("files", [])
    
    system_prompt = (
        f"Kamu adalah {model_name}. User meminta revisi dari jawaban sebelumnya. "
        "Lakukan revisi sesuai permintaan user tanpa mengubah bagian yang tidak diminta. "
        "Berikan jawaban yang sudah direvisi secara LENGKAP (bukan hanya bagian yang berubah). "
        "Pastikan formatnya tetap rapi dengan heading, bullet points, dan markdown."
    )
    
    message = (
        f"Jawaban sebelumnya:\n{original_answer}\n\n"
        f"Permintaan revisi dari user:\n{revision_request}"
    )
    
    return StreamingResponse(
        generate_stream(model_name, system_prompt, message, files),
        media_type="text/event-stream"
    )


# ══════════════════════════════════════════
# Export Endpoints
# ══════════════════════════════════════════

@app.post("/export")
async def export_endpoint(request: Request):
    """Export konsensus ke berbagai format file."""
    data = await request.json()
    content = data.get("content", "")
    format_type = data.get("format", "docx")
    
    if format_type == "docx":
        doc = Document()
        
        # Styling
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Title
        title = doc.add_heading("Consensus AI — Hasil Konsensus", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph("")  # spacing
        
        for line in content.split("\n"):
            line = line.strip()
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("3. "):
                doc.add_paragraph(line[3:], style='List Number')
            elif line:
                doc.add_paragraph(line)
        
        f = io.BytesIO()
        doc.save(f)
        f.seek(0)
        return Response(
            content=f.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=consensus_result.docx"}
        )
        
    elif format_type == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Consensus Result"
        
        # Header style
        header_font = Font(bold=True, size=12)
        
        row_idx = 1
        for line in content.split("\n"):
            if "|" in line and "---" not in line:
                cells = [x.strip() for x in line.split("|") if x.strip()]
                if cells:
                    for col_idx, cell_val in enumerate(cells, 1):
                        cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
                        if row_idx == 1:
                            cell.font = header_font
                            cell.alignment = Alignment(horizontal='center')
                    row_idx += 1
            elif line.strip() and "---" not in line:
                ws.cell(row=row_idx, column=1, value=line.strip())
                row_idx += 1
        
        # Auto-width columns
        for col in ws.columns:
            max_length = 0
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)
                
        f = io.BytesIO()
        wb.save(f)
        f.seek(0)
        return Response(
            content=f.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=consensus_result.xlsx"}
        )
    
    elif format_type == "md":
        # Markdown export
        return Response(
            content=content.encode('utf-8'),
            media_type="text/markdown; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=consensus_result.md"}
        )
    
    elif format_type == "txt":
        return Response(
            content=content.encode('utf-8'),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=consensus_result.txt"}
        )
        
    return {"error": "Format tidak valid. Gunakan: docx, xlsx, md, txt"}


# ══════════════════════════════════════════
# Feedback Endpoints
# ══════════════════════════════════════════

@app.post("/feedback")
async def feedback_endpoint(request: Request):
    """Submit feedback dari user."""
    data = await request.json()
    category = data.get("category", "lainnya")
    message = data.get("message", "")
    
    if not message.strip():
        return {"error": "Pesan feedback tidak boleh kosong"}
    
    if category not in ('bug', 'saran', 'lainnya'):
        category = 'lainnya'
    
    feedback_id = await add_feedback(category, message.strip())
    return {"success": True, "id": feedback_id}


@app.get("/admin/feedbacks")
async def get_feedbacks_endpoint(category: str = None):
    """Get all feedbacks untuk admin panel."""
    feedbacks = await get_all_feedbacks(category)
    return {"feedbacks": feedbacks}


@app.patch("/admin/feedbacks/{feedback_id}")
async def update_feedback_endpoint(feedback_id: int, request: Request):
    """Update status feedback."""
    data = await request.json()
    status = data.get("status", "")
    
    success = await update_feedback_status(feedback_id, status)
    if success:
        return {"success": True}
    return {"error": "Feedback tidak ditemukan atau status tidak valid"}


@app.delete("/admin/feedbacks/{feedback_id}")
async def delete_feedback_endpoint(feedback_id: int):
    """Delete a feedback entry."""
    success = await delete_feedback(feedback_id)
    if success:
        return {"success": True}
    return {"error": "Feedback tidak ditemukan"}


# ══════════════════════════════════════════
# Rating Endpoints
# ══════════════════════════════════════════

@app.post("/rate")
async def rate_endpoint(request: Request):
    """Submit rating untuk jawaban."""
    data = await request.json()
    chat_id = data.get("chat_id", "")
    rating = data.get("rating", 0)
    
    if not (1 <= rating <= 5):
        return {"error": "Rating harus antara 1-5"}
    
    rating_id = await add_rating(chat_id, rating)
    return {"success": True, "id": rating_id}


@app.get("/admin/ratings")
async def get_ratings_endpoint():
    """Get all ratings untuk admin panel."""
    ratings = await get_all_ratings()
    return {"ratings": ratings}


# ══════════════════════════════════════════
# Admin Stats
# ══════════════════════════════════════════

@app.get("/admin/stats")
async def get_stats_endpoint():
    """Get combined statistics untuk admin dashboard."""
    feedback_stats = await get_feedback_stats()
    rating_stats = await get_rating_stats()
    return {
        "feedbacks": feedback_stats,
        "ratings": rating_stats
    }


if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
